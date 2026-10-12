"""Small python-docx toolkit for the manuals (identity "Argus Nocturne", Daylight variant for print)."""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

H1 = RGBColor(0x3F, 0x2F, 0xB0)
H2 = RGBColor(0x5B, 0x47, 0xE5)
INK = RGBColor(0x18, 0x16, 0x30)
MUTED = RGBColor(0x5E, 0x5A, 0x7A)
HEAD_FILL = "2E2B45"
ROW_FILL = "F1EFFA"
NOTE_FILL = "EEF0FF"
WARN_FILL = "FDECEF"
MONO = "Consolas"
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def _shade(properties, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    properties.append(shd)


def _field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, text, end):
        run._r.append(el)


class Manual:
    def __init__(self, title: str, subtitle: str, version: str, header: str, audience: str) -> None:
        self.doc = Document()
        self._figure = 0
        self._setup_page(header)
        self._setup_styles()
        self._cover(title, subtitle, version, audience)

    # -- setup --------------------------------------------------------------------------------------
    def _setup_page(self, header: str) -> None:
        s = self.doc.sections[0]
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin, s.bottom_margin = Cm(2.2), Cm(2.0)
        s.different_first_page_header_footer = True
        hp = s.header.paragraphs[0]
        hp.text = header
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = MUTED
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Strana ")
        r.font.size = Pt(8.5)
        _field(fp, "PAGE", "1")
        r2 = fp.add_run(" z ")
        r2.font.size = Pt(8.5)
        _field(fp, "NUMPAGES", "1")

    def _setup_styles(self) -> None:
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = INK
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.12
        rpr = normal.element.get_or_add_rPr()
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "sk-SK")
        rpr.append(lang)
        for name, size, color, before, after in (("Heading 1", 18, H1, 18, 8), ("Heading 2", 13.5, H2, 14, 4),
                                                 ("Heading 3", 11.5, INK, 10, 3)):
            st = styles[name]
            st.font.name = "Calibri"
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = color
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.keep_with_next = True
        # let Word refresh the table of contents and page fields when the file is opened
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        self.doc.settings.element.append(update)

    def _cover(self, title: str, subtitle: str, version: str, audience: str) -> None:
        d = self.doc
        for _ in range(6):
            d.add_paragraph()
        p = d.add_paragraph()
        r = p.add_run("ArgusGrid")
        r.font.size = Pt(34)
        r.font.bold = True
        r.font.color.rgb = H1
        p = d.add_paragraph()
        r = p.add_run(title)
        r.font.size = Pt(24)
        r.font.color.rgb = INK
        p = d.add_paragraph()
        r = p.add_run(subtitle)
        r.font.size = Pt(13)
        r.font.color.rgb = MUTED
        for _ in range(4):
            d.add_paragraph()
        for label, value in (("Verzia produktu", version), ("Dátum vydania", date.today().strftime("%d. %m. %Y")),
                             ("Určené pre", audience), ("Jazyk", "slovenčina")):
            p = d.add_paragraph()
            a = p.add_run(f"{label}: ")
            a.font.color.rgb = MUTED
            p.add_run(value).bold = True
        d.add_page_break()
        h = d.add_paragraph()
        r = h.add_run("Obsah")
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = H1
        toc = d.add_paragraph()
        _field(toc, 'TOC \\o "1-2" \\h \\z \\u', "Obsah sa doplní po otvorení dokumentu (aktualizujte polia klávesom F9).")
        d.add_page_break()

    # -- content ------------------------------------------------------------------------------------
    def h(self, level: int, text: str) -> None:
        self.doc.add_heading(text, level=level)

    def _runs(self, paragraph, text: str, size: float | None = None) -> None:
        for part in INLINE.split(text):
            if not part:
                continue
            if part.startswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
                if size:
                    run.font.size = Pt(size)
            elif part.startswith("`"):
                run = paragraph.add_run(part[1:-1])
                run.font.name = MONO
                run.font.size = Pt((size or 10.5) - 1)
            else:
                run = paragraph.add_run(part)
                if size:
                    run.font.size = Pt(size)

    def p(self, text: str) -> None:
        self._runs(self.doc.add_paragraph(), text)

    def bullets(self, items: list[str]) -> None:
        for it in items:
            self._runs(self.doc.add_paragraph(style="List Bullet"), it)

    def numbered(self, items: list[str]) -> None:
        """Numbers are written as text so every list restarts at 1 (Word list numbering would continue)."""
        for n, it in enumerate(items, 1):
            par = self.doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.6)
            par.paragraph_format.space_after = Pt(3)
            par.add_run(f"{n}.	")
            par.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            self._runs(par, it)

    def note(self, text: str, warning: bool = False) -> None:
        """Block quote: a note, or a security alert (rules/22 section 4)."""
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        tc_pr = cell._tc.get_or_add_tcPr()
        _shade(tc_pr, WARN_FILL if warning else NOTE_FILL)
        borders = OxmlElement("w:tcBorders")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "24")
        left.set(qn("w:color"), "C42340" if warning else "5B47E5")
        borders.append(left)
        tc_pr.append(borders)
        par = cell.paragraphs[0]
        r = par.add_run("Upozornenie: " if warning else "Poznámka: ")
        r.bold = True
        self._runs(par, text)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def code(self, text: str) -> None:
        for line in text.strip("\n").split("\n"):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.left_indent = Cm(0.5)
            _shade(p._p.get_or_add_pPr(), "F4F3FA")
            r = p.add_run(line or " ")
            r.font.name = MONO
            r.font.size = Pt(9)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def table(self, headers: list[str], rows: list[list[str]], widths_cm: list[float] | None = None,
              mono_cols: tuple[int, ...] = ()) -> None:
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, text in enumerate(headers):
            cell = t.rows[0].cells[i]
            _shade(cell._tc.get_or_add_tcPr(), HEAD_FILL)
            run = cell.paragraphs[0].add_run(text)
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(9.5)
        # repeat the header row on every page
        trpr = t.rows[0]._tr.get_or_add_trPr()
        rep = OxmlElement("w:tblHeader")
        rep.set(qn("w:val"), "true")
        trpr.append(rep)
        for n, row in enumerate(rows):
            cells = t.add_row().cells
            for i, text in enumerate(row):
                if n % 2 == 1:
                    _shade(cells[i]._tc.get_or_add_tcPr(), ROW_FILL)
                par = cells[i].paragraphs[0]
                par.paragraph_format.space_after = Pt(2)
                if i in mono_cols:
                    run = par.add_run(text)
                    run.font.name = MONO
                    run.font.size = Pt(9)
                else:
                    self._runs(par, text, size=9.5)
        if widths_cm:
            for row in t.rows:
                for i, w in enumerate(widths_cm):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure(self, path: Path, caption: str, width_cm: float = 16.0) -> None:
        self._figure += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(f"Obrázok {self._figure}: {caption}")
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
